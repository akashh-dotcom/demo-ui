#!/usr/bin/env python3
"""
Metadata Processor for Book.XML

This module reads metadata from CSV or Excel files and populates the BookInfo
section of the Book.XML file. It supports both .csv and .xls/.xlsx formats.

Expected metadata file naming: metadata.csv or metadata.xls or metadata.xlsx

Expected fields in metadata file:
- ISBN (or isbn)
- Title (or title)
- Subtitle (or subtitle) - optional
- Author (or author, authors) - can be multiple rows or comma-separated
- Publisher (or publisher)
- PublicationDate (or pubdate, date, publication_date)
- Edition (or edition)
- CopyrightYear (or copyright_year, copyright)
- CopyrightHolder (or copyright_holder, holder)

Usage:
    from metadata_processor import populate_bookinfo_from_metadata
    
    success = populate_bookinfo_from_metadata(
        book_xml_path="Book.XML",
        metadata_dir="."  # Directory to search for metadata file
    )
"""

import csv
import logging
import re
from pathlib import Path
from typing import Dict, List, Optional

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

from lxml import etree

logger = logging.getLogger(__name__)


def find_metadata_file(directory: Path) -> Optional[Path]:
    """
    Find metadata file in the given directory.
    
    Looks for:
    - metadata.csv
    - metadata.xls
    - metadata.xlsx
    
    Returns:
        Path to metadata file if found, None otherwise
    """
    candidates = [
        directory / "metadata.csv",
        directory / "metadata.xlsx",
        directory / "metadata.xls",
    ]
    
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            logger.info(f"Found metadata file: {candidate}")
            return candidate
    
    return None


def read_csv_metadata(csv_path: Path) -> Dict[str, any]:
    """
    Read metadata from CSV file.
    
    Expected format (flexible):
    - Two columns: Field, Value
    - Or: Field names in first row, values in second row
    - Field names are case-insensitive
    
    Returns:
        Dictionary with metadata fields
    """
    metadata = {
        'isbn': None,
        'title': None,
        'subtitle': None,
        'authors': [],
        'publisher': None,
        'pubdate': None,
        'edition': None,
        'copyright_year': None,
        'copyright_holder': None
    }
    
    try:
        with open(csv_path, 'r', encoding='utf-8-sig') as f:  # utf-8-sig handles BOM
            reader = csv.reader(f)
            rows = list(reader)
            
            if not rows:
                logger.warning(f"CSV file is empty: {csv_path}")
                return metadata
            
            # Detect format
            # Format 1: Two columns (Field, Value)
            # Format 2: Field names in first row, values in subsequent rows
            
            first_row = rows[0]
            
            # Check if this is a horizontal format (fields in first row)
            if len(rows) >= 2 and len(first_row) > 2:
                # Horizontal format: first row is headers
                headers = [h.strip().lower() for h in first_row]
                values = rows[1] if len(rows) > 1 else []
                
                for idx, header in enumerate(headers):
                    if idx >= len(values):
                        continue
                    value = values[idx].strip()
                    if not value:
                        continue
                    
                    _map_metadata_field(metadata, header, value)
            else:
                # Vertical format: two columns (Field, Value)
                for row in rows:
                    if len(row) < 2:
                        continue
                    field = row[0].strip().lower()
                    value = row[1].strip()
                    
                    if not value:
                        continue
                    
                    _map_metadata_field(metadata, field, value)
        
        logger.info(f"Successfully read metadata from CSV: {csv_path}")
        return metadata
        
    except Exception as e:
        logger.error(f"Error reading CSV metadata: {e}")
        return metadata


def read_excel_metadata(excel_path: Path) -> Dict[str, any]:
    """
    Read metadata from Excel file (.xls or .xlsx).
    
    Expected format (same as CSV):
    - Two columns: Field, Value
    - Or: Field names in first row, values in second row
    
    Returns:
        Dictionary with metadata fields
    """
    metadata = {
        'isbn': None,
        'title': None,
        'subtitle': None,
        'authors': [],
        'publisher': None,
        'pubdate': None,
        'edition': None,
        'copyright_year': None,
        'copyright_holder': None
    }
    
    try:
        suffix = excel_path.suffix.lower()
        
        if suffix == '.xlsx' and HAS_OPENPYXL:
            # Read .xlsx with openpyxl
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb.active
            
            rows = []
            for row in ws.iter_rows(values_only=True):
                # Filter out None values and convert to strings
                row_values = [str(cell).strip() if cell is not None else '' for cell in row]
                if any(row_values):  # Skip empty rows
                    rows.append(row_values)
            
        elif suffix == '.xls' and HAS_XLRD:
            # Read .xls with xlrd
            wb = xlrd.open_workbook(excel_path)
            ws = wb.sheet_by_index(0)
            
            rows = []
            for row_idx in range(ws.nrows):
                row_values = [str(ws.cell_value(row_idx, col_idx)).strip() for col_idx in range(ws.ncols)]
                if any(row_values):  # Skip empty rows
                    rows.append(row_values)
        else:
            logger.error(f"Cannot read Excel file: Missing required library for {suffix}")
            logger.error("Install openpyxl for .xlsx or xlrd for .xls files")
            return metadata
        
        if not rows:
            logger.warning(f"Excel file is empty: {excel_path}")
            return metadata
        
        # Same format detection as CSV
        first_row = rows[0]
        
        if len(rows) >= 2 and len(first_row) > 2:
            # Horizontal format
            headers = [h.strip().lower() for h in first_row]
            values = rows[1] if len(rows) > 1 else []
            
            for idx, header in enumerate(headers):
                if idx >= len(values):
                    continue
                value = values[idx].strip()
                if not value:
                    continue
                
                _map_metadata_field(metadata, header, value)
        else:
            # Vertical format
            for row in rows:
                if len(row) < 2:
                    continue
                field = row[0].strip().lower()
                value = row[1].strip()
                
                if not value:
                    continue
                
                _map_metadata_field(metadata, field, value)
        
        logger.info(f"Successfully read metadata from Excel: {excel_path}")
        return metadata
        
    except Exception as e:
        logger.error(f"Error reading Excel metadata: {e}")
        return metadata


def _map_metadata_field(metadata: Dict, field_name: str, value: str) -> None:
    """
    Map a field name and value to the metadata dictionary.
    
    Handles various field name variations and formats.
    """
    field_name = field_name.lower().strip()
    value = value.strip()
    
    if not value or value.lower() in ('n/a', 'na', 'none', ''):
        return
    
    # ISBN
    if field_name in ('isbn', 'isbn13', 'isbn-13', 'isbn10', 'isbn-10'):
        # Clean ISBN (remove hyphens, spaces)
        isbn_clean = re.sub(r'[^0-9X]', '', value.upper())
        if isbn_clean:
            metadata['isbn'] = isbn_clean
    
    # Title
    elif field_name in ('title', 'book_title', 'booktitle'):
        metadata['title'] = value
    
    # Subtitle
    elif field_name in ('subtitle', 'sub_title', 'book_subtitle'):
        metadata['subtitle'] = value
    
    # Author(s)
    elif field_name in ('author', 'authors', 'author_name', 'by'):
        # Handle multiple authors separated by commas, semicolons, or 'and'
        authors = re.split(r'[,;]|\sand\s', value)
        for author in authors:
            author = author.strip()
            if author and author not in metadata['authors']:
                metadata['authors'].append(author)
    
    # Publisher
    elif field_name in ('publisher', 'publishername', 'publisher_name'):
        metadata['publisher'] = value
    
    # Publication Date
    elif field_name in ('pubdate', 'publication_date', 'date', 'published', 'year', 'publicationdate'):
        metadata['pubdate'] = value
    
    # Edition
    elif field_name in ('edition', 'edition_number'):
        metadata['edition'] = value
    
    # Copyright Year
    elif field_name in ('copyright_year', 'copyrightyear', 'copyright', 'copyright_date'):
        # Extract year if it's in the value
        year_match = re.search(r'\b(19|20)\d{2}\b', value)
        if year_match:
            metadata['copyright_year'] = year_match.group(0)
        else:
            metadata['copyright_year'] = value
    
    # Copyright Holder
    elif field_name in ('copyright_holder', 'copyrightholder', 'holder', 'rights_holder'):
        metadata['copyright_holder'] = value


def create_bookinfo_element(metadata: Dict) -> etree._Element:
    """
    Create a <bookinfo> element from metadata dictionary.
    
    Uses placeholders for missing fields to ensure validation passes.
    """
    bookinfo_elem = etree.Element('bookinfo')
    
    # ISBN (use placeholder if not found)
    isbn_elem = etree.SubElement(bookinfo_elem, 'isbn')
    isbn_elem.text = metadata.get('isbn') or '0000000000000'
    
    # Title (use placeholder if not found)
    title_elem = etree.SubElement(bookinfo_elem, 'title')
    title_elem.text = metadata.get('title') or 'Untitled Book'
    
    # Subtitle (optional - only add if exists)
    if metadata.get('subtitle'):
        subtitle_elem = etree.SubElement(bookinfo_elem, 'subtitle')
        subtitle_elem.text = metadata['subtitle']
    
    # Authors
    authorgroup_elem = etree.SubElement(bookinfo_elem, 'authorgroup')
    authors = metadata.get('authors', [])
    if not authors:
        authors = ['Unknown Author']
    
    for author_name in authors:
        author_elem = etree.SubElement(authorgroup_elem, 'author')
        personname_elem = etree.SubElement(author_elem, 'personname')
        
        # Try to split name into firstname/surname
        parts = author_name.split()
        if len(parts) >= 2:
            firstname_elem = etree.SubElement(personname_elem, 'firstname')
            firstname_elem.text = ' '.join(parts[:-1])
            surname_elem = etree.SubElement(personname_elem, 'surname')
            surname_elem.text = parts[-1]
        else:
            surname_elem = etree.SubElement(personname_elem, 'surname')
            surname_elem.text = author_name
    
    # Publisher
    publisher_elem = etree.SubElement(bookinfo_elem, 'publisher')
    publishername_elem = etree.SubElement(publisher_elem, 'publishername')
    publishername_elem.text = metadata.get('publisher') or 'Unknown Publisher'
    
    # Publication Date
    pubdate_elem = etree.SubElement(bookinfo_elem, 'pubdate')
    pubdate_elem.text = metadata.get('pubdate') or '2024'
    
    # Edition
    edition_elem = etree.SubElement(bookinfo_elem, 'edition')
    edition_elem.text = metadata.get('edition') or '1st Edition'
    
    # Copyright
    copyright_elem = etree.SubElement(bookinfo_elem, 'copyright')
    year_elem = etree.SubElement(copyright_elem, 'year')
    year_elem.text = metadata.get('copyright_year') or metadata.get('pubdate') or '2024'
    holder_elem = etree.SubElement(copyright_elem, 'holder')
    holder_elem.text = metadata.get('copyright_holder') or metadata.get('publisher') or 'Copyright Holder'
    
    return bookinfo_elem


def populate_bookinfo_from_metadata(
    book_xml_path: Path,
    metadata_dir: Path,
    backup: bool = True
) -> bool:
    """
    Read metadata file and populate BookInfo section in Book.XML.
    
    Args:
        book_xml_path: Path to Book.XML file
        metadata_dir: Directory to search for metadata file
        backup: Whether to create a backup of Book.XML before modifying
    
    Returns:
        True if metadata was successfully applied, False otherwise
    """
    # Find metadata file
    metadata_file = find_metadata_file(Path(metadata_dir))
    
    if metadata_file is None:
        print("\n" + "="*70)
        print("⚠ WARNING: Metadata file not available")
        print("="*70)
        print("  Could not find metadata.csv or metadata.xls/metadata.xlsx")
        print("  Book.xml may not be complete - using placeholder values")
        print("  Expected location: " + str(Path(metadata_dir).resolve()))
        print("="*70 + "\n")
        logger.warning("No metadata file found - Book.xml will use placeholder values")
        return False
    
    # Read metadata based on file type
    suffix = metadata_file.suffix.lower()
    
    if suffix == '.csv':
        metadata = read_csv_metadata(metadata_file)
    elif suffix in ('.xls', '.xlsx'):
        if not HAS_OPENPYXL and not HAS_XLRD:
            print("\n" + "="*70)
            print("⚠ ERROR: Cannot read Excel file")
            print("="*70)
            print(f"  Found: {metadata_file.name}")
            print("  Missing required library:")
            print("    - For .xlsx files: pip install openpyxl")
            print("    - For .xls files: pip install xlrd")
            print("  Book.xml will use placeholder values")
            print("="*70 + "\n")
            logger.error("Cannot read Excel metadata - missing required library")
            return False
        metadata = read_excel_metadata(metadata_file)
    else:
        logger.error(f"Unsupported metadata file format: {suffix}")
        return False
    
    # Check if we got any useful metadata
    has_metadata = any([
        metadata.get('isbn'),
        metadata.get('title'),
        metadata.get('authors'),
        metadata.get('publisher'),
        metadata.get('pubdate')
    ])
    
    if not has_metadata:
        print("\n" + "="*70)
        print("⚠ WARNING: Metadata file is empty or invalid")
        print("="*70)
        print(f"  File: {metadata_file.name}")
        print("  No valid metadata fields found")
        print("  Book.xml will use placeholder values")
        print("="*70 + "\n")
        logger.warning("Metadata file contains no valid data")
        return False
    
    # Parse Book.XML
    try:
        parser = etree.XMLParser(remove_blank_text=True, resolve_entities=False)
        tree = etree.parse(str(book_xml_path), parser)
        root = tree.getroot()
    except Exception as e:
        logger.error(f"Failed to parse Book.XML: {e}")
        return False
    
    # Create backup if requested
    if backup:
        backup_path = book_xml_path.with_suffix('.xml.backup')
        try:
            book_xml_path.rename(backup_path)
            logger.info(f"Created backup: {backup_path}")
        except Exception as e:
            logger.warning(f"Could not create backup: {e}")
    
    # Remove existing bookinfo/info elements
    for elem in list(root.findall('.//bookinfo')):
        root.remove(elem)
    for elem in list(root.findall('.//info')):
        root.remove(elem)
    
    # Create new bookinfo element
    bookinfo_elem = create_bookinfo_element(metadata)
    
    # Insert bookinfo at the beginning of root
    root.insert(0, bookinfo_elem)
    
    # Write back to Book.XML
    try:
        tree.write(
            str(book_xml_path),
            encoding='UTF-8',
            xml_declaration=True,
            pretty_print=True
        )
        
        print("\n" + "="*70)
        print("✓ Successfully populated BookInfo from metadata")
        print("="*70)
        print(f"  Metadata file: {metadata_file.name}")
        print(f"  ISBN: {metadata.get('isbn') or '[using placeholder]'}")
        print(f"  Title: {metadata.get('title') or '[using placeholder]'}")
        if metadata.get('authors'):
            print(f"  Author(s): {', '.join(metadata['authors'])}")
        print(f"  Publisher: {metadata.get('publisher') or '[using placeholder]'}")
        print(f"  Date: {metadata.get('pubdate') or '[using placeholder]'}")
        print("="*70 + "\n")
        
        logger.info("Successfully populated BookInfo section from metadata file")
        return True
        
    except Exception as e:
        logger.error(f"Failed to write Book.XML: {e}")
        # Restore backup if write failed
        if backup and backup_path.exists():
            try:
                backup_path.rename(book_xml_path)
                logger.info("Restored backup after write failure")
            except Exception as e2:
                logger.error(f"Failed to restore backup: {e2}")
        return False


if __name__ == "__main__":
    # Example usage
    import sys
    
    logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
    
    if len(sys.argv) < 2:
        print("Usage: python metadata_processor.py <book_xml_path> [metadata_dir]")
        print("\nSearches for metadata.csv or metadata.xls/xlsx in metadata_dir")
        print("and populates the BookInfo section of Book.XML")
        sys.exit(1)
    
    book_xml = Path(sys.argv[1])
    metadata_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else book_xml.parent
    
    if not book_xml.exists():
        print(f"Error: Book.XML not found: {book_xml}")
        sys.exit(1)
    
    success = populate_bookinfo_from_metadata(book_xml, metadata_dir)
    sys.exit(0 if success else 1)
